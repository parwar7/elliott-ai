from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# Reuse the audited document primitives and fixed-width table geometry.
from build_tmc_elliott_report import (
    BLUE, DARK, GREEN, GRAY, INK, MUTED, PALE_BLUE, PALE_GOLD, PALE_GREEN,
    PALE_RED, RED, add_bullet, add_callout, add_table, heading, page,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "DELL_12DA_GETTEX_Elliott_Wave_Full_Report_2026-07-14.docx"


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    p.add_run(text)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.82)
section.bottom_margin = Inches(0.78)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK, 8, 4),
):
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("DELL 12DA | Multi-Degree Elliott Wave Research")
r.font.name = "Calibri"
r.font.size = Pt(9)
r.font.color.rgb = RGBColor.from_string(MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

# Editorial cover.
doc.add_paragraph().paragraph_format.space_after = Pt(44)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("FULL MARKET STRUCTURE REPORT")
r.bold = True
r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
r = p.add_run("DELL 12DA")
r.bold = True
r.font.size = Pt(32)
r.font.color.rgb = RGBColor.from_string(INK)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(16)
r = p.add_run("GETTEX Multi-Degree Elliott Wave Analysis")
r.font.size = Pt(17)
r.font.color.rgb = RGBColor.from_string(DARK)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(24)
r = p.add_run("EUR-denominated TradingView data through 14 July 2026")
r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string(MUTED)
add_callout(doc, "Preferred position", "Cycle I > Primary 3 > Intermediate (3) > Minor 3 > Minute iv active. The June high is Minute iii, not completed Minor 3.", PALE_BLUE)
add_callout(doc, "Current classification", "Minute iv is a strict GETTEX W-X-Y / near-Flat candidate. Its B/X retracement is 87.77%, just below the database's 90% Flat threshold.", PALE_GOLD)
add_callout(doc, "Hard invalidation", "A sustained decline below EUR 225.80 overlaps Minute i territory and invalidates the standard Minor 3 impulse count.", PALE_RED, RED)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Probabilistic technical research; not personalized investment advice")
r.italic = True
r.font.size = Pt(9.5)
r.font.color.rgb = RGBColor.from_string(MUTED)

page(doc, "Executive Assessment")
add_callout(doc, "Bottom line", "DELL remains in a powerful active Minor 3 advance, but the current Minute iv correction is not proven complete. Minute v, Minor 4, Minor 5, Intermediate (4), and Intermediate (5) remain ahead if the impulse survives.", PALE_GREEN, GREEN)
heading(doc, "What changed from the previous report", 2)
add_bullet(doc, "The EUR 415.25 high on 2 June is retained as Minute iii inside Minor 3. It is not labeled as the completion of Minor 3.")
add_bullet(doc, "The current correction is Minute iv, not Minor 4. This follows your corrected degree hierarchy and the two-hour subdivision.")
add_bullet(doc, "GETTEX made its Minor 2 price low at EUR 93.59 on 2 February, later than the NYSE feed's January endpoint. EUR/USD movement changes the strict local extreme.")
add_bullet(doc, "The active correction is classified as W-X-Y under the strict 90% rule because B/X retraced 87.77% of A/W. A near-regular Flat remains the main alternate.")
heading(doc, "Current map", 2)
add_table(doc, ["Level", "Role", "Required interpretation"], [
    ("EUR 415.25", "Minute iii high", "Break suggests B/X extension or early Minute v; recount active correction"),
    ("EUR 402.35", "B/X pivot candidate", "Current resistance and correction decision point"),
    ("EUR 367.55", "First post-B support", "Break helps confirm a developing C/Y decline"),
    ("EUR 309.75", "A/W low", "Likely test if a Flat-style C develops"),
    ("EUR 225.80", "Minute i territory", "Hard standard-impulse invalidation"),
], [1600, 2850, 4910])

page(doc, "Data Integrity and Degree Boundary")
add_callout(doc, "Important data constraint", "GETTEX history before November 2021 contains a corporate-action discontinuity and extremely sparse venue volume. That gap is not a market-generated Elliott wave and must not be counted as one.", PALE_RED, RED)
heading(doc, "How the high-degree count is handled", 2)
add_bullet(doc, "Cycle I and completed Primary 1 are inherited from the corporate-action-normalized master series.")
add_bullet(doc, "GETTEX directly confirms the count from Primary 2 onward, where the series is continuous enough for EUR pivot analysis.")
add_bullet(doc, "Volume is GETTEX venue volume, not consolidated DELL volume. It is useful for relative comparisons within this feed but cannot replace NYSE volume confirmation.")
add_bullet(doc, "The report gives price-rule conclusions priority when venue volume and EWO disagree with a visually clean impulse.")
image = ROOT / "dell_12da_monthly_live.png"
if image.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image), width=Inches(6.35))
    c = doc.add_paragraph("Figure 1. Full GETTEX monthly history. The discontinuities and low early volume limit raw high-degree inference.")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic = True
    c.runs[0].font.size = Pt(9)
    c.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
heading(doc, "Dataset", 2)
add_table(doc, ["Timeframe", "Rows", "Coverage", "Use"], [
    ("Monthly", "125 valid", "Jul 2007-Jul 2026", "Corporate-action and high-degree context"),
    ("Weekly", "301", "Oct 2020-Jul 2026", "Primary 2 and Intermediate structure"),
    ("Daily", "398", "Dec 2024-Jul 2026", "Minor and Minute structure"),
    ("2-hour", "2,215", "Jun 2025-Jul 2026", "Active Minor 3 and Minute iv"),
], [1450, 1300, 2700, 3910])

page(doc, "Cycle and Primary-Degree Map")
add_table(doc, ["Wave", "Price path", "Structure", "Status"], [
    ("Cycle I", "Normalized series", "Large motive advance", "Active"),
    ("Primary 1", "Mar 2020-Feb 2022", "Five-wave motive candidate", "Complete on normalized master series"),
    ("Primary 2", "EUR 49.03 -> 31.70", "ABC zigzag candidate", "Probable complete"),
    ("Primary 3", "EUR 31.70 -> active", "Extended impulse", "Active"),
    ("Primary 4", "Future", "Not formed", "Pending"),
    ("Primary 5", "Future", "Not formed", "Pending"),
], [1500, 2350, 2730, 2780])
heading(doc, "Primary 2 forensic classification", 2)
add_table(doc, ["Leg", "Price path", "Evidence"], [
    ("A", "EUR 49.03 -> 33.45", "Directional first decline"),
    ("B", "EUR 33.45 -> 44.78", "72.72% retracement; below the 90% Flat threshold"),
    ("C", "EUR 44.78 -> 31.70", "Terminal decline; C/A length ratio about 0.84"),
], [1200, 2800, 5360])
doc.add_paragraph("The GETTEX evidence supports the previous regular-zigzag preference. The exact lower-degree five-wave internals remain less reliable than on the primary listing because historical German-session liquidity is thin.")
heading(doc, "Primary 3 position", 2)
add_table(doc, ["Intermediate wave", "Price path", "Status"], [
    ("(1)", "EUR 31.70 -> 166.38", "Probable completed impulse"),
    ("(2)", "EUR 166.38 -> 57.80", "Probable completed ABC zigzag"),
    ("(3)", "EUR 57.80 -> active", "Active; Minor 3 in progress"),
    ("(4)", "Future", "Not formed"),
    ("(5)", "Future", "Not formed"),
], [2100, 3400, 3860])

page(doc, "Primary 3: Completed Intermediate Waves")
heading(doc, "Intermediate (1)", 2)
add_table(doc, ["Minor", "EUR path", "Reading"], [
    ("1", "31.70 -> 64.74", "Initial impulse"),
    ("2", "64.74 -> 57.90", "Corrective pullback"),
    ("3", "57.90 -> 116.60", "Extended motive leg"),
    ("4", "116.60 -> 92.97", "No overlap with Minor 1"),
    ("5", "92.97 -> 166.38", "Terminal advance"),
], [1200, 2900, 5260])
add_callout(doc, "Price-rule verdict", "PASS. Wave 3 is not the shortest, and Wave 4 remained above the Wave 1 high. Wave lengths were 33.04, 58.70, and 73.41 EUR.", PALE_GREEN, GREEN)
heading(doc, "Intermediate (2)", 2)
add_table(doc, ["Leg", "EUR path", "Volume / momentum reading"], [
    ("A", "166.38 -> 77.34", "Strong first decline"),
    ("B", "77.34 -> 137.48", "67.54% recovery; EWO crossed zero"),
    ("C", "137.48 -> 57.80", "C/A ratio 0.895; weekly average volume contracted"),
], [1200, 3000, 5160])
add_bullet(doc, "The 67.54% B retracement is compatible with a zigzag and materially below the normal Flat threshold.")
add_bullet(doc, "The entire correction retraced 80.62% of Intermediate (1), deep but still above the Primary 3 origin.")
add_bullet(doc, "Average weekly GETTEX volume contracted from roughly 39.57K in A to 27.24K in B and 19.54K in C, supporting correction exhaustion.")

page(doc, "Intermediate (3): Minor 1 and Minor 2")
heading(doc, "Minor 1", 2)
add_table(doc, ["Minute", "EUR path", "Status"], [
    ("i", "57.80 -> 77.42", "Complete"),
    ("ii", "77.42 -> 69.65", "Complete correction"),
    ("iii", "69.65 -> 120.62", "Extended advance"),
    ("iv", "120.62 -> 99.50", "Complete; no overlap"),
    ("v", "99.50 -> 147.51", "Complete"),
], [1200, 3000, 5160])
add_callout(doc, "Price-rule verdict", "PASS. Wave 3 is not the shortest and Wave 4 stayed above the Wave 1 high. The five-wave price structure is coherent.", PALE_GREEN, GREEN)
add_callout(doc, "Indicator-engine verdict", "WARNING. Wave iii's absolute daily EWO peak was slightly below Waves i and v on this venue. Under the database's strict momentum rule this does not receive full impulse confirmation, although cumulative volume and price rules support the count.", PALE_GOLD)
heading(doc, "Minor 2", 2)
add_table(doc, ["Leg", "EUR pivot", "Interpretation"], [
    ("A", "147.51 -> 98.89", "Five-down candidate"),
    ("B", "98.89 -> 121.37", "Countertrend recovery"),
    ("C", "121.37 -> 93.59", "Terminal decline; GETTEX endpoint 2 Feb 2026"),
], [1200, 3000, 5160])
doc.add_paragraph("Minor 2 remains a probable ABC zigzag. The GETTEX endpoint is later and slightly lower than the primary NYSE endpoint because the EUR-denominated series contains currency movement in addition to the underlying share movement.")

page(doc, "Active Minor 3: Minute Structure")
add_table(doc, ["Minute wave", "EUR path", "Structure / status"], [
    ("i", "93.59 -> 225.80", "Probable complete five-wave impulse"),
    ("ii", "225.80 -> 196.66", "Probable running Flat"),
    ("iii", "196.66 -> 415.25", "Probable complete extended impulse"),
    ("iv", "415.25 -> active", "W-X-Y preferred; near-Flat alternate"),
    ("v", "Future", "Still expected after iv confirms complete"),
], [1800, 2800, 4760])
heading(doc, "Minute iii internal proof", 2)
add_table(doc, ["Minuette", "EUR path", "Price length"], [
    ("(i)", "196.66 -> 268.15", "71.49"),
    ("(ii)", "268.15 -> 256.50", "Correction"),
    ("(iii)", "256.50 -> 388.15", "131.65"),
    ("(iv)", "388.15 -> 345.55", "No overlap into (i)"),
    ("(v)", "345.55 -> 415.25", "69.70"),
], [1600, 3400, 4360])
add_callout(doc, "Price-rule verdict", "PASS. Minuette (iii) is the longest motive leg, and Minuette (iv) remained well above the Minuette (i) high at EUR 268.15.", PALE_GREEN, GREEN)
add_callout(doc, "EWO qualification", "The two-hour EWO peak lagged into later segments after the explosive May gap. The strict EWO rule warns, but this lag does not invalidate Elliott's hard price rules.", PALE_GOLD)
image = ROOT / "dell_12da_2h_final.png"
if image.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image), width=Inches(6.35))
    c = doc.add_paragraph("Figure 2. GETTEX two-hour chart showing the extended Minute iii advance and overlapping Minute iv correction.")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic = True
    c.runs[0].font.size = Pt(9)
    c.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

page(doc, "Minute iv Forensic Correction Analysis")
add_table(doc, ["Leg", "EUR path", "Evidence", "Status"], [
    ("A / W", "415.25 -> 309.75", "Sharp first correction; EWO crossed zero", "Complete candidate"),
    ("B / X", "309.75 -> 402.35", "87.77% retracement", "Complete candidate, not confirmed"),
    ("C / Y", "Not confirmed", "Needs a clear impulsive or corrective decline", "Pending"),
], [1300, 2200, 3560, 2300])
heading(doc, "Classifier result", 2)
add_callout(doc, "Strict database label", "Complex W-X-Y preferred because the B/X retracement is below 90%. If the next decline is a three-wave structure, this label strengthens.", PALE_BLUE)
add_callout(doc, "Alternate", "Near-regular Flat. The NYSE feed crossed the 90% threshold, while GETTEX reached 87.77%; currency and session timing explain the borderline disagreement.", PALE_GOLD)
heading(doc, "Current lower-degree condition", 2)
add_bullet(doc, "After EUR 402.35, price fell to EUR 367.55 and rebounded as high as EUR 396.50.")
add_bullet(doc, "That movement is not yet a completed five-wave decline. Therefore C/Y cannot be called active with high confidence.")
add_bullet(doc, "A sustained break below EUR 367.55 would improve the case that B/X has ended and C/Y is underway.")
add_bullet(doc, "A break above EUR 415.25 can represent an extended B/X or an early Minute v; the internal subdivision must decide between them.")
heading(doc, "Projection map if C/Y confirms", 2)
add_table(doc, ["Projection", "Level", "Meaning"], [
    ("0.618 x A", "EUR 337.15", "Shallow terminal objective"),
    ("1.000 x A", "EUR 296.85", "A=C equality"),
    ("1.618 x A", "EUR 231.65", "Deep extension near hard overlap risk"),
    ("Hard rule", "EUR 225.80", "Below here, standard Minor 3 impulse is invalid"),
], [2200, 2200, 4960])

page(doc, "Decision Framework and Monitoring")
add_table(doc, ["Scenario", "Required evidence", "Wave implication"], [
    ("Minute iv still active", "Failure below 402.35/415.25; break under 367.55", "C/Y decline develops toward 337.15 or 296.85"),
    ("Minute v begins", "Five-wave advance above 415.25 after a confirmed three-wave corrective low", "Minor 3 resumes upward"),
    ("B/X extends", "Break above 415.25 with corrective, overlapping internals", "Minute iv remains active as expanded Flat/combination"),
    ("Minor 3 invalidated", "Sustained decline below 225.80", "Wave iv overlaps Wave i; full recount required"),
], [1900, 3850, 3610])
heading(doc, "Confirmation sequence", 2)
for text in (
    "Determine whether the move from EUR 402.35 forms five waves down or only three.",
    "Use EUR 367.55 as the first structural support test.",
    "If price declines, compare volume and EWO against the A/W leg; weaker readings would support terminal correction behavior.",
    "Do not declare Minute iv complete until a five-wave rise is followed by a three-wave pullback that holds above the correction low.",
    "Retire the standard Minor 3 impulse immediately if EUR 225.80 is broken on a sustained basis.",
):
    numbered(doc, text)
heading(doc, "Expected future hierarchy", 2)
add_table(doc, ["Order", "Wave still ahead"], [
    ("1", "Minute v of active Minor 3"),
    ("2", "Minor 4 correction"),
    ("3", "Minor 5 advance, completing Intermediate (3)"),
    ("4", "Intermediate (4) correction"),
    ("5", "Intermediate (5) advance, completing Primary 3"),
], [1300, 8060])

page(doc, "Methodology and Final Verdict")
heading(doc, "Database layers applied", 2)
add_bullet(doc, "Elliott hard rules: Wave 3 cannot be shortest; Wave 4 cannot overlap Wave 1 in a standard impulse.")
add_bullet(doc, "EWO: 5-period SMA minus 35-period SMA of median price, where median price is (High + Low) / 2.")
add_bullet(doc, "RSI: Wilder 14-period calculation.")
add_bullet(doc, "Volume: cumulative and average venue volume, with explicit warnings for unequal wave duration and thin GETTEX liquidity.")
add_bullet(doc, "Correction classifier: B retracement below 90% favors W-X-Y; 90%-138.2% supports a Flat when Leg 1 has three subwaves.")
heading(doc, "Final verdict", 2)
add_callout(doc, "Preferred count", "Cycle I > Primary 3 > Intermediate (3) > Minor 3 > Minute iv active.", PALE_GREEN, GREEN)
add_callout(doc, "Current pattern", "Strict GETTEX W-X-Y / near-Flat correction. B/X is a candidate endpoint at EUR 402.35, but C/Y is not yet structurally confirmed.", PALE_BLUE)
add_callout(doc, "Degree correction", "EUR 415.25 is Minute iii. It is not the completion of Minor 3. The previously expected Minute v, Minor 4, and Minor 5 remain ahead.", PALE_GOLD)
heading(doc, "Confidence", 2)
add_table(doc, ["Conclusion", "Confidence"], [
    ("Primary 3 active", "Moderate-high"),
    ("Intermediate (3) active", "High"),
    ("Minor 3 active", "High"),
    ("Minute iii completed at EUR 415.25", "High"),
    ("Minute iv is W-X-Y rather than Flat", "Moderate; feed-sensitive"),
    ("B/X ended at EUR 402.35", "Low-moderate until a decisive decline"),
], [7000, 2360])
add_callout(doc, "Research limitation", "This report describes one probabilistic technical count. It does not forecast guaranteed prices or replace risk management, fundamental analysis, or professional financial advice.", GRAY)

doc.core_properties.title = "DELL 12DA GETTEX Full Elliott Wave Report"
doc.core_properties.subject = "Multi-degree EUR-denominated Elliott Wave analysis"
doc.core_properties.author = "Codex"
doc.save(OUTPUT)
print(OUTPUT)
