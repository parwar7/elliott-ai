from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_tmc_elliott_report import (
    BLUE, DARK, GREEN, GRAY, INK, MUTED, PALE_GOLD, PALE_GREEN, PALE_RED,
    RED, WHITE, add_bullet, add_callout, heading, page, set_cell_text, shade,
    table_geometry,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "DELL_NYSE_Compact_Color_Wave_Report_2026-07-15.docx"

PRIMARY = {
    "P1": {"fill": "D9EAF7", "dark": "2E75B6", "name": "Primary 1"},
    "P2": {"fill": "FCE4D6", "dark": "C65911", "name": "Primary 2"},
    "P3": {"fill": "E2F0D9", "dark": "548235", "name": "Primary 3"},
    "P4": {"fill": "F4CCCC", "dark": "C00000", "name": "Primary 4"},
    "P5": {"fill": "E4DFEC", "dark": "7030A0", "name": "Primary 5"},
}


def colored_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, True, WHITE, 9.2)
        shade(table.rows[0].cells[index], DARK)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for group, values in rows:
        cells = table.add_row().cells
        colors = PRIMARY[group]
        for index, value in enumerate(values):
            set_cell_text(cells[index], value, index == 0, colors["dark"] if index == 0 else INK, 9.2)
            shade(cells[index], colors["fill"])
    table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def legend(doc):
    rows = [(key, (value["name"],)) for key, value in PRIMARY.items()]
    colored_table(doc, ["Color key"], rows, [9360])


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.82)
section.bottom_margin = Inches(0.78)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.34)
section.footer_distance = Inches(0.34)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.7)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08
for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 14, 7),
    ("Heading 2", 13, BLUE, 10, 5),
    ("Heading 3", 11.5, DARK, 7, 3),
):
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("DELL NYSE | Compact Color-Coded Elliott Wave Report")
r.font.name = "Calibri"
r.font.size = Pt(8.8)
r.font.color.rgb = RGBColor.from_string(MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

# Page 1: compact cover and decision.
doc.add_paragraph().paragraph_format.space_after = Pt(24)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(5)
r = p.add_run("DELL")
r.bold = True
r.font.size = Pt(31)
r.font.color.rgb = RGBColor.from_string(INK)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
r = p.add_run("Compact Multi-Degree Elliott Wave Report")
r.font.size = Pt(17)
r.font.color.rgb = RGBColor.from_string(DARK)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(18)
r = p.add_run("NYSE:DELL | Live data reviewed 15 July 2026")
r.font.size = Pt(10.5)
r.font.color.rgb = RGBColor.from_string(MUTED)
add_callout(doc, "Preferred count", "Cycle I > Primary 3 > Intermediate (3) > Minor 3 > Minute iv active.", PALE_GREEN, GREEN)
add_callout(doc, "Current market", "Premarket near $461.10, with a high near $464.77. Minute iv's B wave is extending toward the $469.47 Minute iii high.", PALE_GOLD)
add_callout(doc, "Current decision", "No confirmed trade entry. B is still active; wait for either a regular-session breakout and retest or a five-wave rejection decline.", PALE_RED, RED)
heading(doc, "Color rule", 2)
doc.add_paragraph("Each Primary wave has one color. Every lower-degree wave contained inside that Primary wave keeps the same color.")
legend(doc)

# Page 2: highest-degree map.
page(doc, "Highest-Degree Map")
colored_table(doc, ["Primary wave", "Path", "Form", "Status"], [
    ("P1", ("Primary 1", "$11.76 -> $56.00", "Five-wave motive candidate", "Complete / probable")),
    ("P2", ("Primary 2", "$56.00 -> $30.37", "Regular ABC zigzag", "Complete / probable")),
    ("P3", ("Primary 3", "$30.37 -> active", "Extended impulse", "Active")),
    ("P4", ("Primary 4", "Future", "Correction", "Not formed")),
    ("P5", ("Primary 5", "Future", "Final motive wave", "Not formed")),
], [1800, 2300, 3000, 2260])
heading(doc, "Primary 2 confirmation", 2)
colored_table(doc, ["Wave", "Path", "Evidence"], [
    ("P2", ("A", "$56.00 -> $35.12", "First five-down candidate")),
    ("P2", ("B", "$35.12 -> $47.42", "58.9% retracement; below Flat threshold")),
    ("P2", ("C", "$47.42 -> $30.37", "Five-down candidate; terminal zigzag leg")),
], [1500, 2700, 5160])
add_callout(doc, "Primary 2 verdict", "Regular zigzag remains preferred over W-X-Y because B was shallow and C has a coherent five-wave candidate.", PRIMARY["P2"]["fill"], PRIMARY["P2"]["dark"])
heading(doc, "Primary 3 overview", 2)
colored_table(doc, ["Contained wave", "Path", "Status"], [
    ("P3", ("Intermediate (1)", "$30.37 -> $173.96", "Completed impulse candidate")),
    ("P3", ("Intermediate (2)", "$173.96 -> $64.84", "Completed ABC zigzag candidate")),
    ("P3", ("Intermediate (3)", "$64.84 -> active", "Minor 3 active")),
    ("P3", ("Intermediate (4)", "Future", "Not formed")),
    ("P3", ("Intermediate (5)", "Future", "Not formed")),
], [2300, 3000, 4060])

# Page 3: all current subdivisions stay green.
page(doc, "Inside Active Primary 3")
add_callout(doc, "Reading rule", "All rows on this page are green because every wave shown is contained inside Primary 3.", PRIMARY["P3"]["fill"], PRIMARY["P3"]["dark"])
heading(doc, "Intermediate (3)", 2)
colored_table(doc, ["Degree", "Wave", "Path", "Status"], [
    ("P3", ("Minor", "1", "$64.84 -> $166.83", "Complete five-wave candidate")),
    ("P3", ("Minor", "2", "$166.83 -> $109.88", "Complete ABC correction candidate")),
    ("P3", ("Minor", "3", "$109.88 -> active", "Minute iv active")),
    ("P3", ("Minor", "4", "Future", "Not formed")),
    ("P3", ("Minor", "5", "Future", "Not formed")),
], [1500, 1100, 3000, 3760])
heading(doc, "Active Minor 3", 2)
colored_table(doc, ["Degree", "Wave", "Path", "Status"], [
    ("P3", ("Minute", "i", "$109.88 -> $263.99", "Complete impulse candidate")),
    ("P3", ("Minute", "ii", "$263.99 -> $227.27", "Complete Flat candidate")),
    ("P3", ("Minute", "iii", "$227.27 -> $469.47", "Complete extended impulse")),
    ("P3", ("Minute", "iv", "$469.47 -> active", "B wave extending")),
    ("P3", ("Minute", "v", "Future", "Expected after iv completes")),
], [1500, 1100, 3000, 3760])
heading(doc, "Minute iii internal proof", 2)
colored_table(doc, ["Minuette", "Price path", "Hard-rule result"], [
    ("P3", ("(i)", "$227.27 -> $311.56", "Complete")),
    ("P3", ("(ii)", "$311.56 -> $298.62", "Corrective")),
    ("P3", ("(iii)", "$298.62 -> $443.86", "Extended motive leg")),
    ("P3", ("(iv)", "$443.86 -> $402.35", "No overlap into (i)")),
    ("P3", ("(v)", "$402.35 -> $469.47", "Terminal advance")),
], [1600, 3100, 4660])
add_callout(doc, "Structural result", "Minute iii passes Elliott's hard price rules: Wave (iii) is not shortest and Wave (iv) does not overlap Wave (i).", PRIMARY["P3"]["fill"], PRIMARY["P3"]["dark"])

# Page 4: current structure and visual.
page(doc, "Current Minute iv")
colored_table(doc, ["Leg", "Path", "Classification", "Status"], [
    ("P3", ("A", "$469.47 -> $357.07", "Three-wave decline candidate", "Complete candidate")),
    ("P3", ("B", "$357.07 -> at least $464.77", "About 95.8% retracement", "Still active / unconfirmed")),
    ("P3", ("C", "Not established", "Expected five-wave decline if Flat", "Pending")),
], [1300, 2700, 3100, 2260])
add_callout(doc, "Correction label", "Regular Flat is now preferred because B has retraced more than 90% of A. B can still extend above $469.47 in an expanded Flat.", PRIMARY["P3"]["fill"], PRIMARY["P3"]["dark"])
heading(doc, "Live evidence", 2)
add_bullet(doc, "Premarket price: approximately $461.10; premarket high: approximately $464.77.")
add_bullet(doc, "Two-hour RSI: about 67.9. EWO: about +16.3.")
add_bullet(doc, "Premarket volume is too small to confirm a breakout. Regular-session volume must validate any move above $469.47.")
add_bullet(doc, "The current rise is close to resistance and remains part of B until a clean motive breakout is proven.")
image = ROOT / "dell_nyse_2h_20260715.png"
if image.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image), width=Inches(6.15))
    c = doc.add_paragraph("Figure 1. NYSE:DELL two-hour chart. Minute iii ended at the regular-session $469.47 high; Minute iv remains active.")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic = True
    c.runs[0].font.size = Pt(8.8)
    c.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

# Page 5: compact decision levels.
page(doc, "Confirmation and Decision Levels")
colored_table(doc, ["Level", "Meaning", "Required response"], [
    ("P3", ("$469.47", "Regular-session Minute iii high", "Breakout needs strong volume plus successful retest")),
    ("P3", ("$464.77", "Current premarket B high", "Provisional; can change before/after the open")),
    ("P3", ("$444", "First short-term support", "Loss plus failed retest begins rejection confirmation")),
    ("P3", ("$418.78", "Major lower high/low structure", "Break strongly supports C-wave decline")),
    ("P3", ("$395.30", "0.618 A projection from current B", "First provisional C target")),
    ("P3", ("$357.07", "A-wave low", "Expected test in a normal Flat C")),
    ("P3", ("$352.37", "Provisional A=C target", "Recalculate when B finally ends")),
    ("P3", ("$263.99", "Minute i high", "Hard standard-impulse overlap invalidation")),
], [1700, 3400, 4260])
heading(doc, "Present decision", 2)
add_callout(doc, "No immediate entry", "Do not chase B near $469.47 and do not short while momentum remains positive. Wait for structural confirmation after the regular session opens.", PALE_GOLD)
heading(doc, "Bullish confirmation", 3)
doc.add_paragraph("A regular-session close above $469.47, expanding NYSE volume/EWO, and a three-wave retest that holds above the breakout level.")
heading(doc, "Bearish confirmation", 3)
doc.add_paragraph("A completed five-wave rejection from the final B high, EWO turning negative, RSI losing 45-50, and a failed retest after price breaks $444; confirmation becomes stronger below $418.78.")
heading(doc, "Report conclusion", 2)
add_callout(doc, "Preferred hierarchy", "Primary 3 remains active. Intermediate (3), Minor 3, and Minute iv are active. Minute v, Minor 4, Minor 5, Intermediate (4), and Intermediate (5) remain ahead if the impulse count survives.", PRIMARY["P3"]["fill"], PRIMARY["P3"]["dark"])
p = doc.add_paragraph("Technical research only. Live premarket levels are provisional and this report is not personalized investment advice.")
p.runs[0].italic = True
p.runs[0].font.size = Pt(9)
p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

doc.core_properties.title = "DELL NYSE Compact Color-Coded Elliott Wave Report"
doc.core_properties.subject = "Short multi-degree report with Primary-wave color inheritance"
doc.core_properties.author = "Codex"
doc.save(OUTPUT)
print(OUTPUT)
