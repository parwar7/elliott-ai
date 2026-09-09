from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "DELL_Elliott_Wave_Professional_Report_2026-07-14.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5E6873"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "EAF4EE"
PALE_GOLD = "FFF4D6"
PALE_RED = "FCE8E6"
GREEN = "1F6B45"
GOLD = "7A5A00"
RED = "9B1C1C"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run(run, size=11, bold=False, color="000000", italic=False, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_text(cell, text, *, bold=False, color="000000", size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    set_run(p.add_run(str(text)), size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths_dxa, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        set_cell_shading(header.cells[i], LIGHT_GRAY)
        set_cell_text(header.cells[i], text, bold=True, color=INK, size=font_size)
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_values):
            set_cell_text(row.cells[i], value, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_status_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [1800, 1740, 5820])
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(("Test", "Result", "Interpretation")):
        set_cell_shading(header.cells[i], LIGHT_GRAY)
        set_cell_text(header.cells[i], text, bold=True, color=INK, size=9)
    for test, result, note, tone in rows:
        row = table.add_row()
        prevent_row_split(row)
        set_cell_text(row.cells[0], test, size=8.8)
        set_cell_text(row.cells[1], result, bold=True, color=tone, size=8.8)
        set_cell_text(row.cells[2], note, size=8.8)
        set_cell_shading(row.cells[1], {
            GREEN: PALE_GREEN, GOLD: PALE_GOLD, RED: PALE_RED
        }.get(tone, WHITE))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_callout(doc, label, text, fill=CALLOUT, color=INK):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(f"{label}: "), size=10, bold=True, color=color)
    set_run(p.add_run(text), size=10, color="202124")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        set_run(p.add_run(bold_prefix), bold=True, color=INK)
        set_run(p.add_run(text[len(bold_prefix):]))
    else:
        set_run(p.add_run(text))
    return p


def add_compact_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    for run in p.runs:
        set_run(run, size=10.5)
    if not p.runs:
        set_run(p.add_run(text), size=10.5)
    else:
        p.runs[0].text = text
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    set_run(p.add_run(text), size=8.5, italic=True, color=MUTED)


def add_chart(doc, filename, caption, width=5.5):
    path = ROOT / filename
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(paragraph.add_run("Page "), size=8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.left_indent = Inches(0.5)
    bullet.paragraph_format.first_line_indent = Inches(-0.25)
    bullet.paragraph_format.space_after = Pt(8)
    bullet.paragraph_format.line_spacing = 1.167

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "DELL TECHNOLOGIES | ELLIOTT WAVE TECHNICAL REPORT"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    for run in hp.runs:
        set_run(run, size=8, bold=True, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(58)
    p.paragraph_format.space_after = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("MARKET STRUCTURE REPORT"), size=10, bold=True, color=GOLD)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(9)
    set_run(p.add_run("Dell Technologies Inc."), size=30, bold=True, color=INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    set_run(p.add_run("Multi-Degree Elliott Wave Analysis"), size=17, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(36)
    set_run(p.add_run("NYSE: DELL | Adjusted price series | Analysis date: 14 July 2026"), size=10.5, color=MUTED)

    add_callout(
        doc,
        "Current structural conclusion",
        "Cycle I is active. Primary 3 is active. Within Primary 3, Intermediate (3) remains active and is most plausibly correcting in Minor 4 after a completed Minor 3 advance to $469.47. The live lower-degree pattern is not yet confirmed complete.",
        fill=PALE_BLUE,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(96)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("Prepared from the project Elliott Wave database, TradingView chart evidence, volume, RSI, EWO, Fibonacci relationships, overlap rules, and nested degree constraints."), size=9.5, italic=True, color=MUTED)
    doc.add_page_break()


def build_report():
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    doc.add_heading("Executive Summary", level=1)
    add_body(doc, "The highest defensible count visible on the current public Class C history begins at the March 2020 adjusted low of $11.76. That low is treated as the origin of a local Cycle-degree advance. Earlier movement from the December 2018 listing is intentionally ignored because the left-side structure is incomplete and cannot be assigned confidently under the database's anchor rules.")
    add_status_table(doc, [
        ("Highest active degree", "Cycle I active", "The visible series supports a large motive advance, but Cycle I is not complete.", GREEN),
        ("Primary position", "Primary 3 active", "Primary 1 and Primary 2 are treated as complete; Primary 3 is advancing.", GREEN),
        ("Intermediate position", "Intermediate (3) active", "Intermediate (1) and (2) are complete candidates inside Primary 3.", GREEN),
        ("Current lower degree", "Minor 4 probable", "Price is correcting after a probable completed Minor 3 at $469.47.", GOLD),
        ("Immediate pattern", "Flat candidate", "Minute A down and a roughly 92% Minute B retracement fit a 3-3-5 flat only if 4-hour internals confirm A as a three-wave leg.", GOLD),
    ])
    add_callout(doc, "Decision point", "A confirmed five-wave decline from $460.50 would strengthen the case that Minute C of Minor 4 is active. A sustained break above $469.47 before that proof would force reassessment toward an expanded B, resumed Minor 3, or an early Minor 5 interpretation.", fill=PALE_GOLD, color=GOLD)

    doc.add_heading("Instrument and Data Selection", level=2)
    add_body(doc, "The analysis uses NYSE:DELL, the currently traded Class C common stock. Regular trading began on 28 December 2018 after completion of the Class V transaction. The legacy NASDAQ:DELL series contains older history but ceased trading in 2013 following a cash buyout. Because of the private-company interval and changed capital structure, the two histories are not stitched into one Elliott sequence.")
    add_chart(doc, "dell_monthly.png", "Figure 1. Monthly adjusted NYSE:DELL chart used to anchor the highest visible degree. Screenshot captured from TradingView; the current July candle is incomplete.", width=4.75)

    doc.add_page_break()
    doc.add_heading("Methodology and Confidence Standard", level=1)
    add_body(doc, "The count is a constrained hypothesis, not a visual label exercise. Every proposed child wave must remain inside its parent in both time and price. Motive waves are checked for Wave 2 and Wave 4 overlap, Wave 3 length and strength, internal 5-wave form, and momentum/volume behavior. Corrections are classified from internal wave counts, retracement depth, alternation, Fibonacci relationships, and momentum reset behavior.")
    add_table(doc, ["Status", "Meaning"], [
        ("Confirmed", "All required structural rules are satisfied on the available lower timeframe."),
        ("Probable", "Best-fit count with strong price and indicator evidence, but one or more lower-degree subdivisions remain partially unresolved."),
        ("Unproven internal structure", "The larger pivot sequence is credible, but a required 3-wave or 5-wave internal count has not been demonstrated on sufficiently granular data."),
        ("Invalid", "A hard Elliott rule is violated, such as Wave 3 being the shortest in a standard impulse or Wave 4 overlapping Wave 1 without a diagonal allowance."),
    ], [2100, 7260], font_size=9.3)

    doc.add_heading("Indicator Framework", level=2)
    add_compact_bullet(doc, "EWO = SMA(5) of median price minus SMA(35) of median price, where median price = (High + Low) / 2.")
    add_compact_bullet(doc, "Wave 3 should normally exceed Wave 1 in cumulative volume and absolute EWO peak; Wave 3 cannot be the weakest of Waves 1, 3, and 5.")
    add_compact_bullet(doc, "A healthy Wave 4 should show reduced volume and an EWO reset toward or through zero; absence of the reset means the correction is not yet verified as complete.")
    add_compact_bullet(doc, "A higher Wave 5 price extreme accompanied by lower volume and lower EWO than Wave 3 is treated as exhaustion evidence, not as a standalone reversal signal.")
    add_compact_bullet(doc, "RSI is used for relative momentum confirmation and divergence, while Fibonacci levels are zones of proportionality rather than exact turning-point guarantees.")

    doc.add_heading("Degree Hierarchy", level=2)
    add_table(doc, ["Degree", "Current interpretation", "Status"], [
        ("Cycle", "I: $11.76 (Mar 2020) -> in progress", "Probable active"),
        ("Primary", "1 complete; 2 complete candidate; 3 active", "Probable"),
        ("Intermediate", "Inside Primary 3: (1) complete, (2) complete, (3) active", "Probable"),
        ("Minor", "Inside Intermediate (3): 1, 2, 3 complete candidates; 4 active", "Probable"),
        ("Minute", "Inside Minor 4: A complete candidate, B complete candidate, C not confirmed", "Unproven active"),
    ], [1500, 6060, 1800], font_size=8.9)

    doc.add_page_break()
    doc.add_heading("Cycle I: Primary-Degree Map", level=1)
    add_table(doc, ["Wave", "Price path", "Interpretation", "Confidence"], [
        ("Primary 1", "$11.76 -> $56.00\nMar 2020-Feb 2022", "Completed five-wave motive advance", "Probable"),
        ("Primary 2", "$56.00 -> $30.37\nFeb-Oct 2022", "Regular ABC zigzag preferred; W-X-Y alternate", "Probable / internals unproven"),
        ("Primary 3", "$30.37 -> active\nOct 2022-present", "Extended motive advance; Intermediate (3) active", "Probable active"),
        ("Primary 4", "Future", "Not formed", "Pending"),
        ("Primary 5", "Future", "Not formed", "Pending"),
    ], [1350, 2180, 4030, 1800], font_size=8.7)

    doc.add_heading("Primary 1", level=2)
    add_body(doc, "Primary 1 is counted as a probable five-wave impulse from the March 2020 low to the February 2022 high. The proposed Intermediate pivots are: (1) $11.76 -> $31.17, (2) -> $27.14, (3) -> $48.22, (4) -> $41.75, and (5) -> $56.00. Wave (4) remains above the Wave (1) price territory, preserving the standard impulse rule. Exact Minute subdivisions cannot be recovered confidently from the available older lower-timeframe slice, so this count remains probable rather than fully confirmed.")

    doc.add_heading("Primary 2: Forensic Correction Classification", level=2)
    add_table(doc, ["Leg", "Price path", "Role", "Evidence"], [
        ("A / W", "$56.00 -> $35.12", "First decline", "Directionally sharp; lower-degree five is plausible but not fully proven."),
        ("B / X", "$35.12 -> $47.42", "Countertrend recovery", "Retraced 58.9% of A, much shallower than the typical 90%-138.2% Flat B zone."),
        ("C / Y", "$47.42 -> $30.37", "Terminal decline", "Clean five-wave candidate with no Wave iv / Wave i overlap."),
    ], [1150, 1900, 2050, 4260], font_size=8.7)
    add_body(doc, "The preferred label is a regular zigzag because the B-wave retracement is approximately 58.9%, the correction is directional rather than sideways, and the final decline can be counted as five waves: $47.42 -> $38.86 -> $45.47 -> $31.22 -> $35.38 -> $30.37. Wave iii is not the shortest, and Wave iv remains below Wave i territory. C equals approximately 81.7% of A; the 78.6% projection from B was near $31.01 versus the actual $30.37 low.")
    add_status_table(doc, [
        ("Zigzag test", "Leading", "B is shallow and C has a five-wave candidate, fitting 5-3-5 behavior.", GREEN),
        ("Flat test", "Not supported", "B did not retrace 90%-138.2% of A.", GREEN),
        ("Triangle test", "Rejected", "Only three major legs are present, and a triangle is not valid in a Wave 2 position.", GREEN),
        ("W-X-Y alternate", "Still possible", "If lower-timeframe evidence proves A/W is a three rather than a five, the combination label becomes preferable.", GOLD),
    ])
    add_callout(doc, "Primary 2 verdict", "Probable regular ABC zigzag, with the strict database status 'Unproven Internal Structure' because the internal five waves of A are not fully demonstrated. This is the important distinction between a leading interpretation and a confirmed fact.", fill=PALE_GOLD, color=GOLD)

    doc.add_page_break()
    doc.add_heading("Primary 3: Intermediate-Degree Structure", level=1)
    add_chart(doc, "dell_weekly.png", "Figure 2. Weekly chart supporting the Primary 3 subdivision and the momentum/volume comparisons. Current price and indicator values are provisional while the weekly candle is open.", width=4.5)
    add_table(doc, ["Wave", "Price path", "Structure", "Status"], [
        ("Intermediate (1)", "$30.37 -> $173.96\nOct 2022-May 2024", "Five-wave impulse candidate", "Completed / probable"),
        ("Intermediate (2)", "$173.96 -> $64.84\nMay 2024-Apr 2025", "Regular ABC zigzag candidate", "Completed / probable"),
        ("Intermediate (3)", "$64.84 -> active\nApr 2025-present", "Motive advance; Minor 4 active", "Active / probable"),
        ("Intermediate (4)", "Future", "Not formed", "Pending"),
        ("Intermediate (5)", "Future", "Not formed", "Pending"),
    ], [1850, 2470, 3240, 1800], font_size=8.5)

    doc.add_heading("Intermediate (1): Completed Impulse Candidate", level=2)
    add_table(doc, ["Minor wave", "Price path", "Lower-degree reading"], [
        ("1", "$30.37 -> $69.51", "Probable five: 30.37 -> 42.29 -> 33.80 -> 56.05 -> 51.19 -> 69.51"),
        ("2", "$69.51 -> $61.34", "Three-leg corrective decline"),
        ("3", "$61.34 -> $131.30", "Probable five: 61.34 -> 83.80 -> 77.62 -> 126.38 -> 101.25 -> 131.30"),
        ("4", "$131.30 -> $110.22", "Corrective pullback without overlap into Minor 1"),
        ("5", "$110.22 -> $173.96", "Terminal advance; exact Minute subdivision remains unproven"),
    ], [1350, 2200, 5810], font_size=8.7)
    add_body(doc, "Weekly RSI supports the motive interpretation: Minor 3 recorded the strongest momentum near 89.5, while Minor 5 reached a higher price with lower RSI near 78 before reversal. That is a bearish RSI divergence. Volume did not contract in Minor 5, however, so the stricter volume-divergence reversal rule was not satisfied.")
    add_status_table(doc, [
        ("Price structure", "Passed", "Five Minor pivots and no standard Wave 4 / Wave 1 overlap.", GREEN),
        ("Wave 3 momentum", "Passed", "Minor 3 carried the strongest RSI momentum.", GREEN),
        ("Correction volume", "Warning", "Average volume rose in Minor 2 and Minor 4 rather than drying up.", GOLD),
        ("Wave 5 divergence", "Partial", "RSI diverged, but volume did not contract below Minor 3.", GOLD),
    ])

    doc.add_page_break()
    doc.add_heading("Intermediate (2): Corrective Anatomy", level=1)
    add_table(doc, ["Minor leg", "Price path", "Internal form", "Confirmation"], [
        ("A", "$173.96 -> $84.44", "Five-wave candidate", "Strong directional decline"),
        ("B", "$84.44 -> $143.95", "Three-wave candidate", "Retraced approximately 66.5% of A"),
        ("C", "$143.95 -> $64.84", "Five-wave candidate", "143.95 -> 107.40 -> 122.61 -> 86.45 -> 99.36 -> 64.84"),
    ], [1400, 2050, 2770, 3140], font_size=8.7)
    add_body(doc, "This decline is best classified as a regular zigzag. C is approximately 88% of A, and the entire Intermediate (2) correction retraced roughly 76% of Intermediate (1). Average weekly volume contracted from about 75.53 million in A to 48.80 million in B and 40.60 million in C. That contraction is compatible with a correction rather than an accelerating new bear impulse. RSI also showed a slight bullish divergence: the April 2025 price low at $64.84 was below the March low near $86.45, while RSI was marginally higher.")
    add_status_table(doc, [
        ("Internal sequence", "Passed provisionally", "A and C each have five-wave candidates separated by a three-wave B.", GREEN),
        ("B retracement", "Zigzag-compatible", "Approximately 66.5%, below the normal Flat threshold.", GREEN),
        ("Volume profile", "Passed", "Volume contracted through the correction.", GREEN),
        ("Terminal momentum", "Supportive", "Slight bullish RSI divergence appeared into the final low.", GREEN),
    ])

    doc.add_heading("Intermediate (3): Active Motive Advance", level=1)
    add_table(doc, ["Minor wave", "Price path", "Interpretation", "Status"], [
        ("1", "$64.84 -> $166.83", "Completed five-wave impulse candidate", "Probable complete"),
        ("2", "$166.83 -> $109.88", "Completed ABC zigzag candidate", "Probable complete"),
        ("3", "$109.88 -> $469.47", "Extended five-wave impulse candidate", "Probable complete"),
        ("4", "$469.47 -> active", "Corrective phase; flat candidate", "Active / unconfirmed"),
        ("5", "Future", "Not formed", "Pending"),
    ], [1350, 2150, 3880, 1980], font_size=8.7)

    doc.add_heading("Minor 1 and Minor 2", level=2)
    add_body(doc, "Minor 1 subdivides as $64.84 -> $113.35 -> $104.79 -> $140.55 -> $116.41 -> $166.83. Its fourth wave remains narrowly above the first-wave high, preserving the standard impulse boundary. Minor 2 then forms a probable ABC zigzag: A $166.83 -> $115.88, B -> $140.08, and C -> $109.88.")

    doc.add_heading("Minor 3: Extended Impulse", level=2)
    add_table(doc, ["Minute wave", "Price path", "Reading"], [
        ("i", "$109.88 -> $184.29", "Initial motive leg"),
        ("ii", "$184.29 -> $154.89", "Corrective retracement"),
        ("iii", "$154.89 -> $263.99", "Strong continuation"),
        ("iv", "$263.99 -> $227.27", "Corrective reset"),
        ("v", "$227.27 -> $469.47", "Extended terminal leg; strongest daily EWO"),
    ], [1600, 2600, 5160], font_size=8.8)
    add_body(doc, "The fifth Minute wave extended sharply rather than diverging. Daily RSI reached approximately 91.3 at the June 2026 peak, and daily EWO reached approximately 118.1. Because both momentum and volume expanded into the high, the database's strict Wave 5 volume/EWO divergence rule did not trigger. This matters: the $469.47 high is treated as a structural pivot candidate, not as a classic exhaustion-confirmed top.")
    add_status_table(doc, [
        ("Wave 3 strength", "Passed", "Minor 3 average weekly volume (~43.47M) exceeded Minor 1 (~30.24M).", GREEN),
        ("Wave 3 internals", "Passed provisionally", "Five Minute pivots are visible and Wave iii is not the weakest.", GREEN),
        ("Wave 5 divergence", "Not present", "Price, daily EWO, and volume expanded together into $469.47.", GOLD),
        ("Minor 4 reset", "Not yet passed", "Daily EWO remains positive and has not touched/crossed zero.", GOLD),
    ])

    doc.add_page_break()
    doc.add_heading("Current Focus: Minor 4", level=1)
    add_chart(doc, "dell_daily.png", "Figure 3. Daily view of the late-stage Minor 3 advance and the developing Minor 4 correction. The plotted July 2026 candle and indicator values are incomplete.", width=4.55)
    add_table(doc, ["Minute leg", "Price path", "Current interpretation", "Confidence"], [
        ("A", "$469.47 -> $357.07", "First corrective leg", "Probable complete"),
        ("B", "$357.07 -> $460.50", "Deep countertrend retracement (~92% of A)", "Probable complete"),
        ("C", "$460.50 -> developing?", "Expected terminal five-wave decline if a Flat is unfolding", "Not confirmed"),
    ], [1450, 2300, 3760, 1850], font_size=8.7)
    add_body(doc, "A three-wave Minute A followed by a B retracement between 90% and 138.2% is the database signature for a sideways Flat correction. The observed B retracement is approximately 92%, so the price geometry fits. The classification is still conditional because the 4-hour internals of A and B must be demonstrated as threes and C must develop as a five.")
    add_callout(doc, "Flat scenario", "If Minute C equals Minute A in price length, a rough objective is $348.10. This is a proportional target, not a guaranteed floor. The broader Minor 4 retracement zones are approximately $384.60 (23.6% of Minor 3), $332.10 (38.2%), and $289.70 (50%).", fill=PALE_BLUE)
    add_status_table(doc, [
        ("Leg 1 subwaves", "Unproven", "The Flat label requires Minute A to contain three lower-degree waves on 4-hour data.", GOLD),
        ("Leg 2 depth", "Passed", "Minute B retraced about 92%, inside the Flat diagnostic band.", GREEN),
        ("Volume behavior", "Supportive", "Minor 4 average weekly volume (~32.04M) is below Minor 3 (~43.47M).", GREEN),
        ("EWO zero reset", "Not passed", "Daily EWO remains above zero, so the correction is not verified complete.", GOLD),
        ("Terminal C impulse", "Not confirmed", "A clear five-wave decline from $460.50 is still required.", GOLD),
    ])

    doc.add_heading("Lower-Degree Interpretation", level=2)
    add_body(doc, "On the 4-hour view, price is oscillating below the $469.47 high after the deep B-wave rebound. This behavior can be the early portion of Minute C, but the subdivisions are not mature enough to assign reliable Minuette labels. The correct analytical action is to wait for a complete five-wave decline, then inspect whether the decline terminates near a Fibonacci zone with RSI/EWO divergence and contracting terminal volume.")

    doc.add_page_break()
    doc.add_heading("Scenario Framework and Invalidation", level=1)
    add_table(doc, ["Scenario", "What must happen", "Implication"], [
        ("Base case: Minor 4 Flat", "A and B prove as threes; price forms a five-wave Minute C below $357.07 or into the broader support zone.", "Minor 4 completes, followed by a potential Minor 5 advance inside Intermediate (3)."),
        ("Expanded Flat", "Price first exceeds $469.47 in a B-wave extension, then reverses in five waves.", "Minor 4 remains active; the higher high is not automatically Minor 5."),
        ("Minor 3 still extending", "The market breaks $469.47 with renewed impulse internals and no completed corrective C.", "The $469.47 pivot is relabeled as a lower-degree wave within Minor 3."),
        ("Deeper Minor 4", "Price breaks below $348 and continues through $332 toward $290 while maintaining corrective structure.", "Still compatible with Minor 4, provided it remains above the major structural invalidation boundary."),
        ("Standard Intermediate (3) invalidation", "Minor 4 enters Minor 1 price territory below $166.83 without a valid diagonal context.", "The standard five-wave Intermediate (3) count must be rejected or materially relabeled."),
    ], [1880, 4460, 3020], font_size=8.5)
    add_callout(doc, "Key discipline", "A price level alone does not confirm a wave. The database requires the expected internal structure and indicator behavior at that level. In particular, a move above $469.47 can belong to expanded B, continuing Minor 3, or eventual Minor 5; the subwaves decide the label.", fill=PALE_GOLD, color=GOLD)

    doc.add_heading("Cross-Degree Confirmation Matrix", level=1)
    add_table(doc, ["Sequence", "Structure", "Volume", "RSI / EWO", "Overall"], [
        ("Primary 1", "Probable 5", "Older granular test unavailable", "Broad motive momentum", "Probable"),
        ("Primary 2", "ABC leading; A internals unresolved", "Corrective context", "Bullish RSI divergence into C", "Probable zigzag / unproven"),
        ("Intermediate (1)", "Probable 5", "Correction-volume warnings", "RSI Wave 5 divergence", "Probable impulse"),
        ("Intermediate (2)", "Probable 5-3-5", "Volume contracted A->B->C", "Slight bullish RSI divergence", "Probable zigzag"),
        ("Minor 1 of Int. (3)", "Probable 5", "Avg. ~30.24M", "Motive confirmation", "Probable complete"),
        ("Minor 2 of Int. (3)", "Probable ABC", "Avg. ~32.67M; slight warning", "Corrective", "Probable complete"),
        ("Minor 3 of Int. (3)", "Probable extended 5", "Avg. ~43.47M; strong", "EWO/RSI expanded", "Probable complete"),
        ("Minor 4 of Int. (3)", "Flat candidate", "Avg. ~32.04M; contraction", "EWO zero reset absent", "Active / unconfirmed"),
    ], [2000, 2220, 2000, 2000, 1140], font_size=7.9)

    doc.add_heading("What Is Confirmed vs. What Is Not", level=2)
    add_compact_bullet(doc, "Strongest conclusions: the March 2020 low is the best visible high-degree anchor; the October 2022 and April 2025 lows are major corrective completions; the advance from April 2025 is motive and strong.")
    add_compact_bullet(doc, "Most important open question: whether the current $469.47 -> $357.07 -> $460.50 sequence is definitively 3-3 and therefore the opening of a Flat.")
    add_compact_bullet(doc, "Historical uncertainty: Primary 2 is best labeled ABC, but the first A leg needs lower-timeframe proof of five waves; otherwise W-X-Y remains valid.")
    add_compact_bullet(doc, "No terminal Primary 3 call is justified. Intermediate (3) is not complete while Minor 4 and Minor 5 remain unresolved, and higher Intermediate (4) and (5) have not formed.")

    doc.add_page_break()
    doc.add_heading("Monitoring Checklist", level=1)
    add_table(doc, ["Priority", "Observation", "Confirmation sought"], [
        ("1", "4-hour subdivisions below $460.50", "A clean five-wave decline to identify Minute C."),
        ("2", "Behavior near $357, $348, and $332", "Terminal momentum divergence and declining volume near proportional support."),
        ("3", "Daily EWO", "Touch or cross of zero during Minor 4, followed by an upside turn."),
        ("4", "RSI", "Bullish divergence between lower C-wave price lows, if formed."),
        ("5", "Break above $469.47", "Determine whether internals are corrective (expanded B) or motive (continuing advance)."),
        ("6", "Hard structure boundary", "A decline below $166.83 invalidates the standard Intermediate (3) impulse count unless a diagonal is proven."),
    ], [1050, 3940, 4370], font_size=8.8)

    doc.add_heading("Limitations", level=1)
    add_body(doc, "This report is based on the adjusted NYSE:DELL public history visible from December 2018, with the Elliott anchor placed at March 2020. The July 2026 monthly, weekly, daily, and intraday candles are incomplete. Indicator values can change before candle close. Lower-degree counts become less reliable as the requested subdivision approaches the resolution of the available data. Legacy pre-2013 Dell history was excluded because it represents a different public-company period separated by a private interval and corporate restructuring.")
    add_body(doc, "Elliott Wave analysis is probabilistic. Alternative counts can remain valid until price structure eliminates them. The target zones and invalidation levels in this report are analytical reference points and are not personalized investment advice, a recommendation to trade, or a forecast of guaranteed performance.")

    doc.add_heading("Source Notes", level=1)
    sources = [
        ("Dell Technologies investor relations: stock information", "https://investors.delltechnologies.com/stock-information"),
        ("Dell Technologies: completion of the Class V transaction", "https://investors.delltechnologies.com/news-releases/news-release-details/dell-technologies-completes-class-v-transaction"),
        ("U.S. SEC filing documenting the 2013 legacy DELL transaction", "https://www.sec.gov/Archives/edgar/data/826083/000119312513415302/d610376dsc13e3a.htm"),
        ("TradingView NYSE:DELL symbol page", "https://www.tradingview.com/symbols/NYSE-DELL/"),
    ]
    for label, url in sources:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        add_hyperlink(p, label, url)

    add_callout(doc, "Report status", "Professional analytical report, version 1.0, prepared 14 July 2026. The live count should be updated after a completed daily or weekly pivot materially changes the active Minor 4 structure.", fill=CALLOUT)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
