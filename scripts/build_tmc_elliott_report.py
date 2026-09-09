import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / "tmc_elliott_wave_analysis.json").read_text(encoding="utf-8"))
OUTPUT = ROOT / "TMC_Elliott_Wave_Professional_Report_2026-07-14.docx"

BLUE = "2E74B5"
DARK = "1F4D78"
INK = "17202A"
MUTED = "5E6873"
GRAY = "F2F4F7"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "EAF4EE"
PALE_GOLD = "FFF4D6"
PALE_RED = "FCE8E6"
GREEN = "1F6B45"
RED = "9B1C1C"
WHITE = "FFFFFF"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    props = table._tbl.tblPr
    width = props.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        props.append(width)
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    indent = props.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        props.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    layout = props.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        props.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for col_width in widths:
        node = OxmlElement("w:gridCol")
        node.set(qn("w:w"), str(col_width))
        grid.append(node)
    for row in table.rows:
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            tcw = props.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                props.append(tcw)
            tcw.set(qn("w:w"), str(widths[index]))
            tcw.set(qn("w:type"), "dxa")
            margins = OxmlElement("w:tcMar")
            for side, value in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
                margin = OxmlElement(f"w:{side}")
                margin.set(qn("w:w"), str(value))
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.append(margins)


def set_cell_text(cell, value, bold=False, color=INK, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(str(value))
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, True, WHITE)
        shade(table.rows[0].cells[index], DARK)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            set_cell_text(cells[index], value)
            if row_index % 2:
                shade(cells[index], "F8FAFC")
    table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, label, text, fill=PALE_BLUE, color=DARK):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}: ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.size = Pt(10.5)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(INK)
    table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    p.add_run(text)


def heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def page(doc, title):
    doc.add_page_break()
    heading(doc, title, 1)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.85)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 16, 8), ("Heading 2", 13, BLUE, 12, 6), ("Heading 3", 12, DARK, 8, 4)
):
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = header.add_run("TMC | Elliott Wave Technical Research")
run.font.name = "Calibri"
run.font.size = Pt(9)
run.font.color.rgb = RGBColor.from_string(MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

# Editorial cover.
doc.add_paragraph().paragraph_format.space_after = Pt(48)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("MARKET STRUCTURE REPORT")
r.bold = True
r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
r = p.add_run("TMC")
r.bold = True
r.font.size = Pt(34)
r.font.color.rgb = RGBColor.from_string(INK)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(18)
r = p.add_run("Multi-Degree Elliott Wave Analysis")
r.font.size = Pt(18)
r.font.color.rgb = RGBColor.from_string(DARK)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(26)
r = p.add_run("NASDAQ:TMC | Live TradingView data through 14 July 2026")
r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string(MUTED)
add_callout(doc, "Preferred position", "Primary I active; Intermediate (4) is an active complex W-X-Y correction. Intermediate (5) is expected only after a confirmed corrective low.", PALE_BLUE)
add_callout(doc, "Critical risk", "A sustained break below $3.20 overlaps Intermediate (1) territory and invalidates the standard Primary I impulse count.", PALE_RED, RED)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(22)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Research classification: probabilistic technical analysis, not investment advice")
r.italic = True
r.font.size = Pt(9.5)
r.font.color.rgb = RGBColor.from_string(MUTED)

page(doc, "Executive Assessment")
add_callout(doc, "Bottom line", "TMC is late in a deep Intermediate (4) correction, but the July low is not confirmed as final. Bullish RSI and EWO divergence indicate waning downside momentum; structure confirmation is still required.", PALE_GREEN, GREEN)
heading(doc, "What the count says", 2)
add_bullet(doc, "The December 2022 low at $0.511 is the best operational anchor for Primary I. Earlier SPAC-era price action is treated as listing and collapse context, not as a trustworthy Elliott starting sequence.")
add_bullet(doc, "Intermediate (3), from $0.721 to $11.35, is the clearest completed impulse: all five Minor waves are visible, Wave 3 is not the shortest, and Wave 4 does not overlap Wave 1.")
add_bullet(doc, "The decline from $11.35 is better classified as W-X-Y than as a simple ABC zigzag. Its lower-degree overlap and prolonged sideways middle section are corrective rather than impulsive.")
add_bullet(doc, "The active correction has retraced 70.66% of Intermediate (3). This is deep but legal while price remains above $3.20.")
heading(doc, "Current decision map", 2)
add_table(doc, ["Level", "Meaning", "Analytical response"], [
    ("$3.84", "Current July swing low", "Must hold on a confirmed reversal sequence"),
    ("$4.78-$5.22", "Early resistance band", "Recovery begins to improve the short-term count"),
    ("$6.64", "Lower-degree X-wave high", "Break materially strengthens an Intermediate (4) bottom call"),
    ("$3.20", "Intermediate (1) high", "Sustained break invalidates the standard impulse count"),
], [1500, 3260, 4600])
heading(doc, "Confidence", 2)
add_table(doc, ["Conclusion", "Confidence"], [
    ("Intermediate (3) completed as a five-wave impulse", "High"),
    ("Intermediate (4) is an active W-X-Y", "Moderate-high"),
    ("Primary I remains active", "Moderate"),
    ("The July 2026 low is final", "Low until reversal confirmation"),
], [7000, 2360])

page(doc, "Degree-by-Degree Count")
add_callout(doc, "Degree boundary", "TMC's public price history is too short to confirm a full Cycle or Supercycle structure. Primary is therefore the highest defensible operational degree.", PALE_GOLD)
add_table(doc, ["Degree / wave", "Dates", "Price path", "Status / form"], [
    ("Primary I", "Dec 2022-present", "$0.511 -> active", "Active impulse candidate"),
    ("Intermediate (1)", "Dec 2022-Jul 2023", "$0.511 -> $3.20", "Probable complete impulse"),
    ("Intermediate (2)", "Jul 2023-Dec 2024", "$3.20 -> $0.721", "Probable complex W-X-Y; 92.19% retracement"),
    ("Intermediate (3)", "Dec 2024-Oct 2025", "$0.721 -> $11.35", "Probable complete five-wave impulse"),
    ("Intermediate (4)", "Oct 2025-present", "$11.35 -> $3.84 low", "Active complex W-X-Y"),
    ("Intermediate (5)", "Future", "Not established", "Expected only after (4) confirms complete"),
], [1900, 2100, 2100, 3260])
heading(doc, "Intermediate (2): lower-degree interpretation", 2)
add_table(doc, ["Leg", "Pivot", "Interpretation"], [
    ("W", "$3.20 -> $0.803, Oct 2023", "First corrective decline"),
    ("X", "$0.803 -> $2.07, Mar 2024", "Deep intervening recovery"),
    ("Y", "$2.07 -> $0.721, Dec 2024", "Final complex decline"),
], [1200, 3300, 4860])
doc.add_paragraph("This label is probable rather than confirmed because weekly data cannot resolve every internal subwave. Its duration, repeated reversals, and 92.19% retracement favor a combination over a clean one-directional zigzag.")

page(doc, "Intermediate (3) Impulse Verification")
add_table(doc, ["Minor wave", "Price path", "Weekly avg volume", "EWO peak", "Reading"], [
    ("1", "$0.721 -> $2.55", "15.18M", "0.5996", "Initial advance"),
    ("2", "$2.55 -> $1.57", "14.60M", "0.6969", "Correction"),
    ("3", "$1.57 -> $8.12", "56.32M", "3.0286", "Strongest advancing leg"),
    ("4", "$8.12 -> $4.37", "57.16M", "4.1175", "High-volatility reset"),
    ("5", "$4.37 -> $11.35", "51.07M", "2.7291", "Price high with momentum divergence"),
], [1100, 2050, 1900, 1500, 2810])
heading(doc, "Hard-rule results", 2)
add_bullet(doc, "Wave lengths: Minor 1 = $1.829, Minor 3 = $6.55, Minor 5 = $6.98. Wave 3 is not the shortest, so the Elliott impulse rule passes.")
add_bullet(doc, "Minor 4 bottomed at $4.37, safely above the Minor 1 high at $2.55. The no-overlap rule passes.")
add_bullet(doc, "Minor 3 greatly exceeded Minor 1 in average volume and EWO momentum. The database's Wave 3 strength test passes.")
add_bullet(doc, "Minor 5 exceeded the Minor 3 price high while its EWO peak fell from 3.0286 to 2.7291 and average volume fell from 56.32M to 51.07M. This is a valid price/volume/EWO exhaustion divergence.")
add_callout(doc, "Qualification", "Minor 4 volume did not dry up. It was a high-volume correction, so the sequence does not receive a perfect 'healthy correction' score even though the price rules and Wave 5 divergence strongly support completion.", PALE_GOLD)

page(doc, "Intermediate (4) Forensic Correction Count")
add_table(doc, ["Leg", "Date / price", "Retracement or form", "Status"], [
    ("Start", "13 Oct 2025, $11.35", "Intermediate (3) peak", "Confirmed pivot"),
    ("W", "17 Nov 2025, $4.75", "First corrective family", "Probable complete"),
    ("X", "22 Jan 2026, $10.05", "80.30% recovery of W", "Probable complete"),
    ("Y", "22 Jan 2026-present", "$10.05 -> $3.84 low", "Active / unconfirmed"),
], [1200, 2350, 3510, 2300])
heading(doc, "Why W-X-Y is preferred over ABC", 2)
add_bullet(doc, "The X leg retraced 80.30% of W. This is below the 90% threshold normally used to elevate a regular or expanded Flat interpretation.")
add_bullet(doc, "Inside Y, the visible path is approximately w: $10.05 -> $3.93, x: $3.93 -> $6.64, y: $6.64 -> $3.84. That is a corrective three-leg family.")
add_bullet(doc, "Trying to force the January-to-July decline into a five-wave C creates fourth-wave overlap in the most natural lower-degree pivots. That weakens the simple zigzag case.")
add_bullet(doc, "The entire correction is prolonged and overlapping, behavior that is structurally consistent with a combination.")
heading(doc, "Lower-degree active Y", 2)
add_table(doc, ["Subleg", "Price path", "Indicator character"], [
    ("w", "$10.05 -> $3.93", "Strongest downside momentum; RSI reached 25.11"),
    ("x", "$3.93 -> $6.64", "Complex recovery; EWO crossed positive"),
    ("y", "$6.64 -> $3.84", "Marginal new low with weaker downside momentum"),
], [1300, 2650, 5410])

page(doc, "Momentum, Volume, and Reversal Evidence")
add_callout(doc, "Bullish divergence", "Price fell from the March low of $3.93 to a lower July low of $3.84, while RSI rose from 25.11 to 31.07 and EWO improved from -1.3179 to -0.9173.", PALE_GREEN, GREEN)
add_table(doc, ["Test", "March 30 low", "July 8 low", "Result"], [
    ("Price", "$3.93", "$3.84", "Lower low"),
    ("Daily RSI(14)", "25.11", "31.07", "Bullish divergence"),
    ("Daily EWO", "-1.3179", "-0.9173", "Bullish divergence"),
    ("Daily volume", "9.25M", "6.24M", "Reduced effort at new low"),
], [2200, 2100, 2100, 2960])
heading(doc, "What this confirms", 2)
add_bullet(doc, "Downside momentum is weaker at the new price low.")
add_bullet(doc, "The final y leg has lower average daily volume than the earlier w decline, consistent with selling pressure drying up.")
add_bullet(doc, "The correction may be approaching a terminal region.")
heading(doc, "What this does not confirm", 2)
add_bullet(doc, "Divergence alone does not prove that Intermediate (4) has ended.")
add_bullet(doc, "A new low can still form while RSI and EWO divergence persists.")
add_bullet(doc, "The count needs a completed lower-degree five-wave advance and a corrective pullback that holds above the terminal low.")

page(doc, "Scenario Framework")
add_table(doc, ["Scenario", "Required evidence", "Wave implication", "Confidence"], [
    ("Bullish reversal", "Five up from $3.84; pullback holds; break above $6.64", "Intermediate (4) likely complete; Intermediate (5) begins", "Unconfirmed"),
    ("Correction extends", "Failure below $4.78-$5.22 and another low above $3.20", "Y continues or subdivides into a more complex ending", "Live"),
    ("Primary impulse fails", "Sustained break below $3.20", "Standard Primary I impulse invalid; full high-degree recount", "Hard trigger"),
], [1900, 3100, 3260, 1100])
heading(doc, "Operational confirmation sequence", 2)
for text in (
    "Track whether the rebound from $3.84 forms five clear waves on the four-hour chart.",
    "Require the next pullback to form three waves and remain above $3.84.",
    "Treat $4.78 and $5.22 as early resistance tests, not final confirmation.",
    "Use a decisive break above $6.64 as the stronger structural signal that the final y leg has ended.",
    "Immediately retire the standard impulse count on sustained trade below $3.20.",
):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.add_run(text)
add_callout(doc, "Reference target", "A simple A=C projection would point near $3.45, but this is only a reference because W-X-Y is the preferred correction label. It is not a required target.", PALE_GOLD)

page(doc, "Source Chart and Method")
image_path = ROOT / "tmc_initial.png"
if image_path.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Inches(6.35))
    caption = doc.add_paragraph("Figure 1. NASDAQ:TMC full-history TradingView chart used for the high-degree anchor review.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(9)
    caption.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
heading(doc, "Data and calculations", 2)
add_bullet(doc, "Source: live TradingView Table View for NASDAQ:TMC, captured 14 July 2026.")
add_bullet(doc, "Coverage used: 300 weekly rows, 333 daily rows, and 333 four-hour rows. The intraday file begins in March 2026, so earlier lower-degree labels rely on daily and weekly evidence.")
add_bullet(doc, "EWO is calculated as the 5-period SMA minus the 35-period SMA of median price, where median price is (High + Low) / 2.")
add_bullet(doc, "RSI uses Wilder's 14-period method. Wave volume is evaluated using both cumulative volume and average volume because waves have unequal duration.")
add_bullet(doc, "Labels are probabilistic and must be updated when pivots or invalidation levels are breached.")
add_callout(doc, "Important", "This report is technical research and does not provide personalized investment advice, expected return, or a guarantee of future price behavior.", GRAY)

doc.core_properties.title = "TMC Multi-Degree Elliott Wave Analysis"
doc.core_properties.subject = "NASDAQ:TMC technical market structure report"
doc.core_properties.author = "Codex"
if __name__ == "__main__":
    doc.save(OUTPUT)
    print(OUTPUT)
