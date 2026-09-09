from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn


path = Path(__file__).resolve().parents[1] / "DELL_Elliott_Wave_Professional_Report_2026-07-14.docx"
document = Document(path)
issues = []

for section_number, section in enumerate(document.sections, 1):
    values = (
        section.page_width,
        section.page_height,
        section.top_margin,
        section.right_margin,
        section.bottom_margin,
        section.left_margin,
    )
    expected = (7772400, 10058400, 914400, 914400, 914400, 914400)
    if values != expected:
        issues.append(f"section {section_number} geometry {values}")

for table_number, table in enumerate(document.tables, 1):
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    table_indent = properties.find(qn("w:tblInd"))
    if table_width is None or table_width.get(qn("w:w")) != "9360":
        issues.append(f"table {table_number} width")
    if table_indent is None or table_indent.get(qn("w:w")) != "120":
        issues.append(f"table {table_number} indent")
    grid = [int(column.get(qn("w:w"))) for column in table._tbl.tblGrid]
    if sum(grid) != 9360:
        issues.append(f"table {table_number} grid {sum(grid)}")
    for row_number, row in enumerate(table.rows, 1):
        if len(row.cells) != len(grid):
            issues.append(f"table {table_number} row {row_number} cell count")

with ZipFile(path) as archive:
    bad_member = archive.testzip()
    document_xml = archive.read("word/document.xml").decode("utf-8")
    media = [name for name in archive.namelist() if name.startswith("word/media/")]
    relationships = archive.read("word/_rels/document.xml.rels").decode("utf-8")

word_count = sum(len(paragraph.text.split()) for paragraph in document.paragraphs)
word_count += sum(
    len(cell.text.split())
    for table in document.tables
    for row in table.rows
    for cell in row.cells
)

print("zip_integrity=", bad_member or "OK")
print(
    "paragraphs=", len(document.paragraphs),
    "tables=", len(document.tables),
    "inline_shapes=", len(document.inline_shapes),
)
print("manual_page_breaks=", document_xml.count('w:type="page"'))
print("media_files=", len(media), media)
print("external_hyperlinks=", relationships.count('TargetMode="External"'))
print("word_count=", word_count)
print("geometry_issues=", issues or "NONE")
